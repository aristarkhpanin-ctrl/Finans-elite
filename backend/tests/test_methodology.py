"""Карта методических трактовок (SPEC §22, фаза D2).

Проверяется не арифметика — её держит golden-master, — а **обещания карты**: она не
подтверждает трактовки за человека, не выдаёт наличие поля за задействованность, не
молчит там, где пункт не задействован, и не расходится со списком в спецификации.
"""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

from calc_core import ENGINE_VERSION, run
from calc_core.methodology import RESOLUTIONS, methodology_map
from calc_core.samples import build_sample_project, build_showcase_project

SPEC = Path(__file__).resolve().parents[2] / "docs" / "CALC-ENGINE-SPEC.md"


def _map(model):
    return methodology_map(model, run(model))


def _by_id(report, choice_id):
    return next(c for c in report.choices if c.id == choice_id)


# --- Связь с методикой ---

def test_the_map_lists_exactly_the_open_questions_of_the_spec():
    """Список открытых вопросов живёт в двух местах — и они не должны разойтись.

    Раньше он был только в документе, и код о нём не знал вовсе: пункт, закрытый в SPEC,
    остался бы в карте (и наоборот) до тех пор, пока кто-нибудь не прочёл бы оба текста
    подряд. Теперь расхождение — падающий тест.
    """
    section = SPEC.read_text().split("## 22. Открытые методические вопросы")[1]
    numbered = {int(m) for m in re.findall(r"^(\d+)\.\s+\*\*", section, re.MULTILINE)}
    assert numbered, "в SPEC §22 не нашлось ни одного пункта — изменился формат раздела?"

    report = _map(build_sample_project())
    assert {c.number for c in report.choices} == numbered
    # И номера не повторяются: два пункта под одним номером означали бы, что один из них
    # никто не проверит.
    assert len({c.number for c in report.choices}) == len(report.choices)


def test_every_choice_names_what_is_still_open():
    """Пункт без открытого вопроса — это закрытый пункт, и ему не место в §22."""
    for choice in _map(build_showcase_project()).choices:
        assert choice.open_question, f"{choice.id}: не назван открытый вопрос"
        assert choice.spec.startswith("SPEC §"), f"{choice.id}: не назван раздел методики"
        assert choice.chosen, f"{choice.id}: не сказано, что выбрано"


# --- Что карта утверждает и чего не утверждает ---

def test_the_map_does_not_confirm_anything_by_itself():
    """Подтверждение трактовки — профессиональное суждение человека на реальных
    проектах. Карта, объявляющая себя подтверждением, — ровно та ложь, ради которой
    версия ядра и держится предварительной."""
    report = _map(build_sample_project())
    assert report.confirmed is False
    assert "не подтверждены" in report.note


def test_the_engine_version_stays_preliminary_while_questions_are_open():
    """Пока в §22 есть хоть один открытый пункт, `engine_version` — `0.x`.

    Фиксация `1.0` это утверждение «трактовки подтверждены», и сделать его вправе
    человек, а не очередной коммит. Тест не даёт цифре уехать молча.
    """
    report = _map(build_sample_project())
    assert report.choices, "открытых вопросов не осталось — фиксация 1.0 обсуждается людьми"
    assert ENGINE_VERSION.startswith("0."), (
        "версия ядра перестала быть предварительной, а открытые методические вопросы "
        "в SPEC §22 остались: либо закрыть их, либо вернуть 0.x")


# --- Задействованность ---

def test_a_choice_is_engaged_by_numbers_not_by_the_presence_of_a_field():
    """«Курсовая разница учитывается» и «курсовая разница здесь равна нулю» — разные
    ответы: первый, сказанный про вторую модель, отправляет проверяющего искать то,
    чего нет."""
    plain = _map(build_sample_project())
    fx = _by_id(plain, "fx.revaluation")
    assert fx.engaged is False
    assert "нет" in fx.silent_because and fx.evidence["i25_total"] == "0"

    rich = _by_id(_map(build_showcase_project()), "fx.revaluation")
    assert rich.engaged is True and Decimal(rich.evidence["i25_total"]) > 0


def test_a_silent_choice_names_the_reason_instead_of_keeping_quiet():
    for choice in _map(build_sample_project()).choices:
        if not choice.engaged:
            assert choice.silent_because, f"{choice.id}: молчит без причины"
        else:
            # И наоборот: причина молчания у задействованного пункта — противоречие.
            assert not choice.silent_because, f"{choice.id}: назван и живым, и молчащим"


def test_vat_is_judged_by_its_own_lines_not_by_total_taxes():
    """C12 несёт и налог на прибыль, и имущество: при выключенном НДС непустая C12
    читалась бы как «НДС всё-таки есть»."""
    off = _by_id(_map(build_sample_project()), "vat.basis")
    assert off.engaged is False and off.evidence["vat_rate"] == "0"
    assert "vat_receivable_b7" in off.evidence


def test_the_shipment_and_payment_bases_are_described_differently():
    model = build_sample_project()
    model.settings.vat_rate = Decimal("0.20")
    shipment = _by_id(methodology_map(model, run(model)), "vat.basis")
    assert shipment.engaged and "по отгрузке" in shipment.chosen.lower()

    model.settings.vat_basis = "payment"
    payment = _by_id(methodology_map(model, run(model)), "vat.basis")
    assert "по оплате" in payment.chosen.lower()
    assert payment.chosen != shipment.chosen


def test_auto_financing_is_silent_when_switched_off():
    model = build_sample_project()
    off = _by_id(methodology_map(model, run(model)), "financing.auto")
    assert off.engaged is False and "выключен" in off.silent_because

    model.financing.auto_financing.enabled = True
    on = _by_id(methodology_map(model, run(model)), "financing.auto")
    assert on.engaged is True and not on.silent_because


def test_i24_names_where_it_came_from():
    """Число без источника нечего проверять: «41 тыс. ₽ за счёт прибыли» — откуда?"""
    engaged = _by_id(_map(build_showcase_project()), "profit.i24")
    assert engaged.engaged is True and engaged.evidence["sources"]


# --- Слой не трогает расчёт ---

def test_the_map_is_a_reader_and_changes_nothing():
    """Как ревью плана: чистая функция над готовым результатом. Карта, способная
    поправить методику, была бы второй методикой."""
    model = build_showcase_project()
    result = run(model)
    before = model.model_dump_json()
    numbers = {code: list(series) for code, series in result.income.lines.items()}

    methodology_map(model, result)

    assert model.model_dump_json() == before
    assert {code: list(s) for code, s in result.income.lines.items()} == numbers
    # И в самом результате карты нет: golden-снимок движка она не двигает.
    assert not hasattr(result, "methodology")


# --- Выходы: API и документ ---

def test_the_api_answers_by_the_project(client, register):
    """Тот же вопрос, но про конкретный проект: карта считается по его модели."""
    headers = register()
    model = client.get("/api/v1/sample").json()
    pid = client.post("/api/v1/projects", json={"name": "П", "model": model},
                      headers=headers).json()["id"]

    body = client.get(f"/api/v1/projects/{pid}/methodology", headers=headers).json()
    assert body["confirmed"] is False and "не подтверждены" in body["note"]
    assert body["engine_version"].startswith("0.")
    assert len(body["choices"]) == 8
    assert body["engaged_count"] == sum(1 for c in body["choices"] if c["engaged"])
    silent = [c for c in body["choices"] if not c["engaged"]]
    assert silent and all(c["silent_because"] for c in silent)


def test_the_document_prints_the_assumptions_and_names_the_omitted(client, register):
    """Бизнес-план читают бухгалтер, аудитор и банк: им важно не только число, но и
    трактовка, по которой оно получено. Раньше она жила только в спецификации движка."""
    from io import BytesIO

    from docx import Document

    headers = register()
    model = client.get("/api/v1/sample").json()
    pid = client.post("/api/v1/projects", json={"name": "П", "model": model},
                      headers=headers).json()["id"]

    content = client.get(f"/api/v1/projects/{pid}/business-plan.docx", headers=headers).content
    text = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
    assert "Методические допущения расчёта" in text
    assert "не подтверждены" in text
    # Усечённый список называет свою неполноту — как и в заключении.
    assert "из 8" in text


# --- Классификация открытых вопросов (как каждый закрывается) ---

def test_every_open_question_says_how_it_closes():
    """Восемь пунктов весили одинаково — и это была неправда.

    «Какие строки усреднять в коэффициентах» и «начисляется ли НДС с аванса» стояли рядом
    как равные, хотя у второго есть норма, а у первого её нет и быть не может. Из-за
    этого работа человека выглядела как восемь открытых вопросов вместо четырёх.
    """
    report = _map(build_showcase_project())
    for choice in report.choices:
        assert choice.resolution in RESOLUTIONS, f"{choice.id}: неизвестный способ закрытия"
    # Каждое состояние кем-то занято: свободное «на будущее» обещает работу, которой нет.
    assert {c.resolution for c in report.choices} == set(RESOLUTIONS)


def test_a_citable_question_carries_the_norm_and_a_judgement_one_does_not():
    """Смысл класса — в том, что он несёт с собой. «Закрывается ссылкой» без ссылки —
    то же самое обещание, что и раньше, только с новым словом."""
    for choice in _map(build_showcase_project()).choices:
        if choice.resolution == "citable":
            assert choice.proposed_basis, f"{choice.id}: норма не названа"
            assert "ст." in choice.proposed_basis or "ПБУ" in choice.proposed_basis
        else:
            assert not choice.proposed_basis, (
                f"{choice.id}: основание названо там, где нормы нет — "
                "это выдало бы суждение за норму")


def test_the_classification_confirms_nothing():
    """Ровно та ложь, ради которой версия и держится предварительной: «предложена норма»
    читается как «уже согласовано», если рядом не написано обратное."""
    report = _map(build_sample_project())
    assert report.confirmed is False
    assert "предложение" in report.classification_note
    assert ENGINE_VERSION.startswith("0.")


def test_fewer_questions_wait_for_a_human_than_the_list_suggested():
    """Число, ради которого классификация и затевалась."""
    report = _map(build_showcase_project())
    assert 0 < len(report.needs_human) < len(report.choices)
    assert all(c.resolution == "judgement" for c in report.needs_human)


# --- Расхождения с нормой: отдельно от «ещё обсуждается» ---

def test_the_loss_carryforward_divergence_is_named_where_it_bites():
    """Пул убытков покрывает базу целиком, а п. 2.1 ст. 283 НК разрешает не больше
    половины. Это не трактовка: у вопроса есть ответ, и он другой."""
    from calc_core.templates import INDUSTRY_TEMPLATES

    losing = _by_id(_map(INDUSTRY_TEMPLATES["cafe"].build()), "tax.loss_carryforward")
    assert losing.divergence and "50%" in losing.divergence
    assert "283" in losing.proposed_basis

    # У модели без переноса убытков расхождения нет — и предупреждения тоже.
    profitable = _by_id(_map(INDUSTRY_TEMPLATES["retail"].build()), "tax.loss_carryforward")
    assert profitable.divergence == ""


def test_the_advance_vat_divergence_needs_all_three_conditions():
    """НДС с аванса расходится с п. 1 ст. 167 НК только там, где есть что начислять:
    НДС включён, режим «по отгрузке» и в сбыте действительно есть предоплата."""
    from calc_core.models.common import VatBasis
    from calc_core.templates import INDUSTRY_TEMPLATES

    prepaid = INDUSTRY_TEMPLATES["saas"].build()             # предоплата 100%
    assert _by_id(_map(prepaid), "vat.basis").divergence

    # «По оплате» — НДС и так идёт за деньгами, расхождения нет.
    by_payment = INDUSTRY_TEMPLATES["saas"].build()
    by_payment.settings.vat_basis = VatBasis.PAYMENT
    assert _by_id(_map(by_payment), "vat.basis").divergence == ""

    # Без НДС начислять нечего.
    no_vat = INDUSTRY_TEMPLATES["saas"].build()
    no_vat.settings.vat_rate = Decimal(0)
    assert _by_id(_map(no_vat), "vat.basis").divergence == ""

    # Без предоплаты — тоже: у кофейни деньги в момент отгрузки.
    assert _by_id(_map(INDUSTRY_TEMPLATES["cafe"].build()), "vat.basis").divergence == ""


def test_a_divergence_is_named_not_fixed():
    """Правка расхождения меняет числа и требует своего бампа версии с ревью
    golden-диффа. Карта — читающий слой: она о нём только рассказывает."""
    from calc_core.templates import INDUSTRY_TEMPLATES

    model = INDUSTRY_TEMPLATES["cafe"].build()
    before = run(model)
    _map(model)                                   # разбор ничего не меняет
    after = run(model)
    assert before.income["I22"] == after.income["I22"]
    assert before.income["I27"] == after.income["I27"]


def test_the_document_prints_the_divergence_under_its_own_heading():
    """В общем списке «открыто к сверке» расхождение читается как «ещё обсуждается».
    Читателю бизнес-плана — банку, аудитору — важно именно оно."""
    from datetime import date as _date
    from io import BytesIO

    from docx import Document

    from app.docgen import build_business_plan_docx
    from calc_core.review import ReviewContext, run_review
    from calc_core.review.opinion import build_opinion
    from calc_core.templates import INDUSTRY_TEMPLATES

    model = INDUSTRY_TEMPLATES["cafe"].build()
    result = run(model)
    opinion = build_opinion(run_review(ReviewContext(model=model, result=result)), result)
    doc = Document(BytesIO(build_business_plan_docx(
        model, result, opinion, project_name="Кофейня", today=_date(2026, 7, 1))))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Где расчёт расходится с нормой" in text
    assert "283" in text and "50%" in text
    assert "предложение" in text                  # и оговорка едет рядом
