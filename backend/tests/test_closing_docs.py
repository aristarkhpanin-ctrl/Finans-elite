"""Закрывающие документы: счёт на оплату и акт (пакет G, G6).

Проверяются обещания:

* сумма прописью правильна во всех формах (1 рубль · 2–4 рубля · 11–14 рублей · тысяча
  женского рода);
* реквизиты продавца **не готовы — документ не формируется**, и отказ говорит выход;
  опечатка в счёте продавца (ключ по БИК) ловится, а не печатается;
* реквизиты покупателя — у организации, ИНН проверяется тем же кодом, что у дела;
* документ — **снимок**: смена реквизитов после составления не меняет перепечатку;
* акт — на каждый успешный платёж, **датой окончания периода**, один на платёж; пока
  реквизитов нет — ждёт и потом составляется той же датой;
* платёж без записанного периода назван пробелом, а не угадан;
* документы **переживают организацию** и читаются платформой после её удаления.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from docx import Document

from app import closing_docs, crud, scheduler
from app.closing_docs import valid_account, valid_corr_account
from app.closing_docx import vat_line
from app.database import as_tenant
from app.db_models import AuditLogEntry, BillingDocument, Payment
from app.main import app
from app.money_words import rubles_in_words

BIK = "044525225"
CORR = "30101810400000000225"          # известный верный корр. счёт для этого БИК
SELLER_INN, SELLER_KPP = "7707083893", "773601001"
BUYER_INN, BUYER_KPP = "7736050003", "772801001"
IP_INN = "500100732259"


def _with_key(bik: str, account: str) -> str:
    """Подобрать контрольную цифру (9-ю) расчётного счёта под БИК."""
    for digit in "0123456789":
        candidate = account[:8] + digit + account[9:]
        if valid_account(bik, candidate):
            return candidate
    raise AssertionError("ключ не подбирается")


ACCOUNT = _with_key(BIK, "40702810000000001234")


# --- Сумма прописью ---

@pytest.mark.parametrize("rub, words", [
    (0, "Ноль рублей 00 копеек"),
    (1, "Один рубль 00 копеек"),
    (2, "Два рубля 00 копеек"),
    (5, "Пять рублей 00 копеек"),
    (11, "Одиннадцать рублей 00 копеек"),
    (12, "Двенадцать рублей 00 копеек"),
    (14, "Четырнадцать рублей 00 копеек"),
    (21, "Двадцать один рубль 00 копеек"),
    (22, "Двадцать два рубля 00 копеек"),
    (111, "Сто одиннадцать рублей 00 копеек"),
    (1000, "Одна тысяча рублей 00 копеек"),
    (2000, "Две тысячи рублей 00 копеек"),
    (2900, "Две тысячи девятьсот рублей 00 копеек"),
    (11000, "Одиннадцать тысяч рублей 00 копеек"),
    (21000, "Двадцать одна тысяча рублей 00 копеек"),
    (31320, "Тридцать одна тысяча триста двадцать рублей 00 копеек"),
    (1_000_000, "Один миллион рублей 00 копеек"),
    (2_000_000, "Два миллиона рублей 00 копеек"),
    (5_000_000, "Пять миллионов рублей 00 копеек"),
    (1_001_001, "Один миллион одна тысяча один рубль 00 копеек"),
])
def test_amount_in_words(rub, words):
    assert rubles_in_words(rub) == words


@pytest.mark.parametrize("kop, tail", [(1, "01 копейка"), (2, "02 копейки"),
                                       (11, "11 копеек"), (21, "21 копейка")])
def test_kopecks_agree_too(kop, tail):
    assert rubles_in_words(3, kop).endswith(tail)


def test_out_of_range_is_refused_not_guessed():
    with pytest.raises(ValueError):
        rubles_in_words(-1)
    with pytest.raises(ValueError):
        rubles_in_words(5, 100)


# --- Счета по ключу Банка России ---

def test_a_known_correspondent_account_passes_and_a_typo_does_not():
    assert valid_corr_account(BIK, CORR)
    assert not valid_corr_account(BIK, CORR[:-1] + "6")


def test_a_settlement_account_typo_is_caught():
    assert valid_account(BIK, ACCOUNT)
    broken = ACCOUNT[:-1] + str((int(ACCOUNT[-1]) + 1) % 10)
    assert not valid_account(BIK, broken)


# --- Продавец ---

@pytest.fixture
def seller(monkeypatch):
    for var, value in {"SELLER_NAME": "ООО «Финанс-Элит»", "SELLER_INN": SELLER_INN,
                       "SELLER_KPP": SELLER_KPP, "SELLER_ADDRESS": "Москва, ул. Тверская, 1",
                       "SELLER_BANK": "ПАО Сбербанк", "SELLER_BIK": BIK,
                       "SELLER_ACCOUNT": ACCOUNT, "SELLER_CORR_ACCOUNT": CORR,
                       "SELLER_DIRECTOR": "Петров П. П."}.items():
        monkeypatch.setenv(var, value)
    monkeypatch.delenv("SELLER_ACCOUNTANT", raising=False)
    monkeypatch.delenv("SELLER_VAT_RATE", raising=False)
    return monkeypatch


def test_without_seller_requisites_every_missing_one_is_named(monkeypatch):
    for var, _ in closing_docs.SELLER_ENV.values():
        monkeypatch.delenv(var, raising=False)
    problems = closing_docs.seller_problems()
    for key in closing_docs.SELLER_REQUIRED:
        assert any(closing_docs.SELLER_ENV[key][0] in p for p in problems), key


def test_complete_seller_requisites_have_no_problems(seller):
    assert closing_docs.seller_problems() == []


@pytest.mark.parametrize("var, value, fragment", [
    ("SELLER_INN", "7707083894", "контрольной цифры"),
    ("SELLER_KPP", "", "нужен КПП"),
    ("SELLER_ACCOUNT", ACCOUNT[:-1] + str((int(ACCOUNT[-1]) + 1) % 10), "SELLER_ACCOUNT"),
    ("SELLER_BIK", "04452522", "9 цифр"),
    ("SELLER_VAT_RATE", "двадцать", "SELLER_VAT_RATE"),
])
def test_a_typo_in_seller_requisites_blocks_documents(seller, var, value, fragment):
    """По счёту уходят деньги: опечатка в нём не печатается, а останавливает документ."""
    seller.setenv(var, value)
    assert any(fragment in p for p in closing_docs.seller_problems())


def test_a_sole_proprietor_has_no_kpp(seller):
    seller.setenv("SELLER_INN", IP_INN)
    assert any("КПП не бывает" in p for p in closing_docs.seller_problems())
    seller.setenv("SELLER_KPP", "")
    assert closing_docs.seller_problems() == []


# --- Счёт через API ---

def _org(client, headers) -> str:
    return client.get("/api/v1/organizations", headers=headers).json()[0]["id"]


def _requisites(client, org_id, headers, **over):
    body = {"legal_name": "ООО «Клиент»", "inn": BUYER_INN, "kpp": BUYER_KPP,
            "legal_address": "Москва, ул. Новая, 5", **over}
    return client.put(f"/api/v1/organizations/{org_id}/billing/requisites", json=body,
                      headers=headers)


def _invoice(client, org_id, headers, plan="team", months=1):
    return client.post(f"/api/v1/organizations/{org_id}/billing/invoices",
                       json={"plan_code": plan, "months": months}, headers=headers)


def _docx_text(client, org_id, doc_id, headers) -> str:
    r = client.get(f"/api/v1/organizations/{org_id}/billing/documents/{doc_id}/docx",
                   headers=headers)
    assert r.status_code == 200
    doc = Document(BytesIO(r.content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_no_seller_requisites_no_invoice_and_the_way_out_is_named(client, register,
                                                                  db_session, monkeypatch):
    monkeypatch.delenv("SELLER_NAME", raising=False)
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    r = _invoice(client, org_id, headers)
    assert r.status_code == 409 and "Напишите платформе" in r.json()["detail"]
    assert db_session.query(BillingDocument).count() == 0


def test_buyer_requisites_are_required_and_named(client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    r = _invoice(client, org_id, headers)
    assert r.status_code == 409 and "реквизиты организации" in r.json()["detail"]
    saved = _requisites(client, org_id, headers, inn="7736 050004").json()
    assert saved["inn"] == "7736050004"                 # пробелы убраны
    assert any("контрольной цифры" in p for p in saved["problems"])   # но опечатка названа
    assert _invoice(client, org_id, headers).status_code == 409


def test_an_invoice_is_numbered_printed_and_logged(client, register, seller, db_session):
    headers = register(email="owner@e.ru")
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    first = _invoice(client, org_id, headers, months=12).json()
    second = _invoice(client, org_id, headers).json()
    assert (first["number"], second["number"]) == (1, 2)
    assert first["amount_rub"] == 2900 * 12 and first["title"].startswith("Счёт № 1 от")

    text = _docx_text(client, org_id, first["id"], headers)
    assert "Счёт на оплату № 1" in text and "Тридцать четыре тысячи восемьсот" in text
    assert f"ООО «Клиент», ИНН {BUYER_INN}, КПП {BUYER_KPP}" in text
    assert ACCOUNT in text and CORR in text and "Без налога (НДС)" in text
    assert "Главный бухгалтер" not in text              # не задан — не печатается
    with as_tenant(db_session, org_id):
        actions = {e.action for e in db_session.query(AuditLogEntry).filter_by(
            organization_id=org_id)}
    assert {"org.requisites_update", "billing.invoice",
            "billing.document_download"} <= actions


def test_a_reprint_is_the_same_document_after_requisites_change(client, register, seller):
    """Документ — снимок: перепечатка через год не должна стать другим документом под
    тем же номером."""
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    doc = _invoice(client, org_id, headers).json()
    _requisites(client, org_id, headers, legal_name="ООО «Переименованный»")
    seller.setenv("SELLER_DIRECTOR", "Сидоров С. С.")
    text = _docx_text(client, org_id, doc["id"], headers)
    assert "ООО «Клиент»" in text and "Переименованный" not in text
    assert "Петров П. П." in text


def test_a_price_on_request_plan_has_no_invoice(client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    assert _invoice(client, org_id, headers, plan="audit_corp").status_code == 409


def test_another_organization_cannot_read_the_document(client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    doc = _invoice(client, org_id, headers).json()
    stranger = register(email="other@e.ru", org="Чужая")
    other_org = _org(client, stranger)
    r = client.get(f"/api/v1/organizations/{other_org}/billing/documents/{doc['id']}/docx",
                   headers=stranger)
    assert r.status_code == 404


def test_vat_is_named_or_its_absence_is():
    assert vat_line(2900, None) == ("Без налога (НДС)", "—")
    assert vat_line(2900, 22) == ("В том числе НДС (22%)", "522,95")


# --- Акты ---

NOW = datetime(2026, 9, 20, 12, 0, tzinfo=timezone.utc)


def _paid(db, org_id: str, *, start: datetime = NOW, months: int = 1) -> Payment:
    """Оплата, прошедшая через ту же дверь, что у провайдеров: период записан."""
    from app import billing
    from app.plans import PLANS

    payment = crud.create_payment(db, org_id, "team", 2900 * months, months=months)
    billing.activate_paid_plan(db, org_id, PLANS["team"], paid_at=start, months=months,
                               payment=payment)
    crud.mark_payment(db, payment, "succeeded")
    return payment


def test_an_act_is_dated_by_the_end_of_the_paid_period(client, register, seller,
                                                        db_session):
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    payment = _paid(db_session, org_id)
    end = payment.period_end
    assert end is not None

    upcoming = client.get(f"/api/v1/organizations/{org_id}/billing/documents",
                          headers=headers).json()["upcoming"]
    assert upcoming[0]["state"] == "scheduled"

    assert scheduler.issue_acts(db_session, NOW + timedelta(days=5)).issued == 0
    assert scheduler.issue_acts(db_session, NOW + timedelta(days=31)).issued == 1
    assert scheduler.issue_acts(db_session, NOW + timedelta(days=32)).issued == 0  # один
    act = db_session.query(BillingDocument).filter_by(kind="act").one()
    assert act.doc_date == (NOW + timedelta(days=30)).date()
    text = _docx_text(client, org_id, act.id, headers)
    assert "Акт № 1" in text and "за период с 20.09.2026 по 20.10.2026" in text
    assert "Заказчик ____________________ " in text   # имя заказчика не выдумано


def test_an_act_waits_for_requisites_and_keeps_its_date(client, register, seller,
                                                        db_session):
    headers = register()
    org_id = _org(client, headers)
    _paid(db_session, org_id)
    run = scheduler.issue_acts(db_session, NOW + timedelta(days=31))
    assert run.issued == 0 and run.blocked == 1
    state = client.get(f"/api/v1/organizations/{org_id}/billing/documents",
                       headers=headers).json()["upcoming"][0]
    assert state["state"] == "blocked" and "реквизиты организации" in state["reason"]

    _requisites(client, org_id, headers)
    scheduler.issue_acts(db_session, NOW + timedelta(days=60))
    act = db_session.query(BillingDocument).filter_by(kind="act").one()
    assert act.doc_date == (NOW + timedelta(days=30)).date()   # дата периода, не заполнения


def test_a_payment_without_a_recorded_period_is_named_not_guessed(client, register,
                                                                  seller, db_session):
    headers = register()
    org_id = _org(client, headers)
    payment = crud.create_payment(db_session, org_id, "team", 2900)
    crud.mark_payment(db_session, payment, "succeeded")
    upcoming = client.get(f"/api/v1/organizations/{org_id}/billing/documents",
                          headers=headers).json()["upcoming"]
    assert upcoming[0]["state"] == "no_period" and "по запросу" in upcoming[0]["reason"]
    assert scheduler.issue_acts(db_session, NOW + timedelta(days=400)).issued == 0


def test_an_invoice_paid_by_transfer_gets_its_months_and_period(client, register,
                                                                db_session):
    """Найдено в G6: назначение оператором заводило платёж без числа месяцев — платёж за
    квартал выглядел месячным, а акт по нему был бы не на тот срок."""
    headers = register()
    org_id = _org(client, headers)
    staff = register(email="staff@e.ru", org="Платформа")
    crud.set_staff(db_session, crud.get_user_by_email(db_session, "staff@e.ru"),
                   is_staff=True)
    client.post(f"/api/v1/admin/organizations/{org_id}/subscription",
                json={"plan_code": "team", "months": 3}, headers=staff)
    payment = db_session.query(Payment).filter_by(organization_id=org_id).one()
    assert payment.months == 3 and payment.period_end is not None


def test_the_nightly_run_leaves_a_trace(db_session):
    scheduler.issue_acts(db_session, NOW)
    assert scheduler.last_runs(db_session)["acts"] is not None


# --- Документы переживают клиента ---

def test_documents_outlive_the_organization_and_the_platform_can_read_them(
        client, register, seller, db_session):
    headers = register(email="owner@e.ru", org="Уходящая")
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    doc = _invoice(client, org_id, headers).json()
    preview = client.get(f"/api/v1/organizations/{org_id}/delete-preview",
                         headers=headers).json()
    assert any("Счета и акты" in k for k in preview["kept"])
    r = client.request("DELETE", f"/api/v1/organizations/{org_id}",
                       json={"password": "secret123"}, headers=headers)
    assert r.status_code == 200

    assert db_session.query(BillingDocument).filter_by(id=doc["id"]).count() == 1
    support = register(email="support@e.ru", org="Платформа")
    crud.set_staff(db_session, crud.get_user_by_email(db_session, "support@e.ru"),
                   is_staff=True, role="support")
    listed = client.get("/api/v1/admin/billing-documents", params={"org_id": org_id},
                        headers=support).json()
    assert listed[0]["organization_exists"] is False
    assert listed[0]["buyer_name"] == "ООО «Клиент»"
    r = client.get(f"/api/v1/admin/billing-documents/{doc['id']}/docx", headers=support)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_the_export_lists_documents(client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    _invoice(client, org_id, headers)
    export = client.get(f"/api/v1/organizations/{org_id}/export", headers=headers).json()
    assert export["документы"][0]["документ"].startswith("Счёт № 1")


def test_the_documents_screen_names_what_is_not_issued(client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    body = client.get(f"/api/v1/organizations/{org_id}/billing/documents",
                      headers=headers).json()
    assert body["seller_ready"] is True and "УПД" in body["not_issued"]


def test_numbering_restarts_each_year(db_session, client, register, seller):
    headers = register()
    org_id = _org(client, headers)
    _requisites(client, org_id, headers)
    org = crud.get_organization(db_session, org_id)
    from app.plans import PLANS
    a = closing_docs.create_invoice(db_session, org, PLANS["team"], 1, requested_by="x",
                                    today=date(2026, 12, 31))
    b = closing_docs.create_invoice(db_session, org, PLANS["team"], 1, requested_by="x",
                                    today=date(2027, 1, 1))
    assert (a.number, b.number) == (1, 1)


def test_app_is_importable_with_documents_routes():
    paths = app.openapi()["paths"]
    assert "/api/v1/organizations/{org_id}/billing/invoices" in paths
    assert "/api/v1/admin/billing-documents" in paths
