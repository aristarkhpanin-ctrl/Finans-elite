"""Абонентская база: приток и отток вместо ручного объёма (SPEC §5).

До 0.9.42 подписку в модели можно было задать только рядом объёмов — то есть **вручную
посчитанной базой**. Отток при этом не выражался никак, и шаблон «Сервис по подписке»
вынужден был писать первым же пунктом: база только растёт, без оттока подписная модель
красива всегда.

Проверяется рекуррента, её стыки с остальной моделью (цена, предоплата, рецептура, старт
продукта, инфляция) и два отказа: подмена ручного объёма **называется**, а нулевой отток
получает предупреждение ревью — он допустим, но не нейтрален.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

from calc_core import run
from calc_core.models import (
    BomLine,
    Company,
    Financing,
    Material,
    OperatingPlan,
    PaymentTerms,
    Product,
    ProjectHeader,
    ProjectModel,
    ProjectSettings,
    SalesLine,
    StartingBalance,
    Subscription,
)
from calc_core.money import quantize
from calc_core.review import run_review
from calc_core.review.types import ReviewContext

D = Decimal


def _model(sub: Subscription | None = None, *, n=12, volume=None, price=D(1000),
           payment=None, bom=False, start_month=None, **settings) -> ProjectModel:
    opts = dict(discount_rate_annual=D("0.15"), profit_tax_rate=D("0.20"),
                property_tax_rate=D("0"), vat_rate=D("0"))
    opts.update(settings)
    op = OperatingPlan(
        materials=[Material(id="m1", name="Хостинг", unit_price=D(100))] if bom else [],
        products=[Product(id="p1", name="Подписка",
                          bom=[BomLine(material_id="m1", qty_per_unit=D(1))] if bom else [])],
        sales=[SalesLine(product_id="p1", price=[price] * n, subscription=sub,
                         volume=volume or [], start_month=start_month,
                         payment=payment or PaymentTerms())],
    )
    return ProjectModel(
        header=ProjectHeader(name="Подписка", start_date=date(2026, 1, 1), duration_months=n),
        settings=ProjectSettings(**opts),
        company=Company(starting_balance=StartingBalance()),
        operating_plan=op,
        financing=Financing(common_shares=D(100)),
    )


def _base(r) -> list[Decimal]:
    return r.subscription_base[0].base


# --- Рекуррента ---

def test_the_base_grows_by_arrivals_and_shrinks_by_churn():
    """База считается на бумаге: 100 на старте, 10 приходят, 10% уходят.

    мес.0: 100 − 10 + 10 = 100 · мес.1: 100 − 10 + 10 = 100 — приток ровно покрывает
    выбытие, и база стоит на месте. Это и есть равновесие, к которому идёт любая подписка.
    """
    sub = Subscription(starting_base=D(100), new_per_month=[D(10)] * 6,
                       churn_monthly=D("0.1"))
    r = run(_model(sub, n=6))
    assert [quantize(v) for v in _base(r)] == [quantize(D(100))] * 6


def test_churn_is_taken_before_arrivals_so_a_new_subscriber_does_not_leave_at_once():
    """Порядок внутри месяца — не придирка: при быстром росте он меняет базу заметно.

    50 на старте, 100 приходят, 20% уходят. Правильно: 50 − 10 + 100 = 140. Если считать
    отток после притока, вышло бы (50 + 100) · 0,8 = 120 — уволили бы пятую часть тех,
    кто подписался вчера.
    """
    sub = Subscription(starting_base=D(50), new_per_month=[D(100)],
                       churn_monthly=D("0.2"))
    r = run(_model(sub, n=1))
    assert quantize(_base(r)[0]) == quantize(D(140))
    assert quantize(r.subscription_base[0].churned[0]) == quantize(D(10))


def test_the_base_converges_to_arrivals_over_churn():
    """Потолок базы — приток делённый на отток, и модель обязана к нему идти, а не расти
    вечно. Именно этого и не умела модель до появления оттока."""
    sub = Subscription(starting_base=D(0), new_per_month=[D(60)] * 120,
                       churn_monthly=D("0.05"))
    r = run(_model(sub, n=120))
    ceiling = D(60) / D("0.05")                      # 1200
    assert _base(r)[-1] < ceiling
    assert _base(r)[-1] > ceiling * D("0.99")        # подошла вплотную, но не перешла
    assert all(_base(r)[t] <= _base(r)[t + 1] for t in range(119))   # монотонно снизу


def test_zero_churn_is_plain_accumulation():
    """Ноль оттока — допустимое значение: база = стартовая плюс весь приток."""
    sub = Subscription(starting_base=D(10), new_per_month=[D(5)] * 4, churn_monthly=D(0))
    r = run(_model(sub, n=4))
    assert [quantize(v) for v in _base(r)] == [quantize(D(x)) for x in (15, 20, 25, 30)]


def test_full_churn_leaves_only_the_arrivals_of_the_month():
    """Отток 100% — вырожденный край: остаются ровно пришедшие в этом месяце."""
    sub = Subscription(starting_base=D(500), new_per_month=[D(7)] * 3, churn_monthly=D(1))
    r = run(_model(sub, n=3))
    assert [quantize(v) for v in _base(r)] == [quantize(D(7))] * 3


def test_a_short_arrivals_row_is_padded_not_repeated():
    """Ряд короче горизонта — это «дальше никто не приходит», а не «повторить последнее»:
    домыслить приток значило бы дорисовать пользователю продажи, которых он не закладывал."""
    sub = Subscription(starting_base=D(0), new_per_month=[D(100)], churn_monthly=D("0.5"))
    r = run(_model(sub, n=3))
    assert [quantize(v) for v in _base(r)] == [quantize(D(x)) for x in (100, 50, 25)]


# --- Стыки с остальной моделью ---

def test_the_base_becomes_the_sales_volume():
    """Ниже по конвейеру подписки не существует: выручка = база × цена."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(10)] * 3, churn_monthly=D(0))
    r = run(_model(sub, n=3, price=D(2000)))
    assert [quantize(v) for v in r.income["I1"]] == [
        quantize(b * D(2000)) for b in _base(r)]


def test_prepayment_works_on_the_derived_volume():
    """Предоплата — самая частая черта подписки, и она обязана видеть выведенный объём:
    деньги приходят на месяц раньше выручки."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(0)] * 4, churn_monthly=D(0))
    terms = PaymentTerms(prepayment_share=D(1), advance_lead_months=1)
    r = run(_model(sub, n=4, price=D(1000), payment=terms))
    assert r.cashflow["C1"][0] > r.income["I1"][0]       # получено больше, чем начислено
    assert r.balance["B24"][0] > 0                       # разница — аванс


def test_a_recipe_follows_the_derived_base():
    """Хостинг «на абонента» через рецептуру следует за базой сам — второго ряда, который
    надо править вместе с притоком, не заводится."""
    sub = Subscription(starting_base=D(10), new_per_month=[D(10)] * 3, churn_monthly=D(0))
    r = run(_model(sub, n=3, bom=True))
    assert [quantize(v) for v in r.income["I5"]] == [
        quantize(b * D(100)) for b in _base(r)]


def test_the_product_start_gates_arrivals_not_the_base():
    """Старт продукта обнуляет **приток**, а не выведенную базу.

    Иначе абоненты копились бы до старта и вываливались одним скачком в первый же месяц
    продаж — базу, которую ещё некому обслуживать, модель показала бы как достижение.
    """
    sub = Subscription(starting_base=D(1000), new_per_month=[D(50)] * 6,
                       churn_monthly=D("0.1"))
    r = run(_model(sub, n=6, start_month=3))
    base = _base(r)
    assert [quantize(v) for v in base[:3]] == [quantize(D(0))] * 3   # до старта пусто
    assert quantize(base[3]) == quantize(D(50))                      # ровно первый приток
    assert quantize(base[4]) == quantize(D(50) * D("0.9") + D(50))
    assert quantize(r.income["I1"][2]) == 0


def test_price_inflation_applies_to_the_derived_revenue():
    """Индексация цен работает поверх выведенного объёма — как у обычной строки сбыта."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(0)] * 13, churn_monthly=D(0))
    r = run(_model(sub, n=13, inflation_sales=D("0.12")))
    assert r.income["I1"][12] > r.income["I1"][0]        # тот же объём, цена выше


def test_the_balance_still_converges():
    sub = Subscription(starting_base=D(40), new_per_month=[D(9)] * 12, churn_monthly=D("0.04"))
    r = run(_model(sub, n=12, bom=True))
    assert [quantize(v) for v in r.balance["B20"]] == [
        quantize(v) for v in r.balance["B34"]]


def test_a_model_without_subscriptions_is_untouched():
    """Правило инертно: без подписки ни базы в результате, ни предупреждений."""
    r = run(_model(None, n=3, volume=[D(5)] * 3))
    assert r.subscription_base == [] and r.warnings == []
    assert [quantize(v) for v in r.income["I1"]] == [quantize(D(5000))] * 3


# --- Отказы, которые называют себя ---

def test_a_manual_volume_next_to_a_subscription_is_named_not_silently_dropped():
    """Два источника объёма молча разошлись бы, и расчёт пошёл бы по тому, о котором
    пользователь забыл. В расчёт идёт база, но подмена **сказана вслух**."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(0)] * 3, churn_monthly=D(0))
    r = run(_model(sub, n=3, volume=[D(999)] * 3))
    assert [quantize(v) for v in _base(r)] == [quantize(D(100))] * 3
    assert len(r.warnings) == 1
    assert "Подписка" in r.warnings[0] and "не используется" in r.warnings[0]


def test_an_empty_manual_volume_is_not_a_conflict():
    """Оставшийся от прежней правки пустой ряд — не конфликт, и шуметь о нём нельзя:
    предупреждение, которое видно всегда, перестают читать."""
    sub = Subscription(starting_base=D(1), new_per_month=[D(0)], churn_monthly=D(0))
    assert run(_model(sub, n=1, volume=[D(0), D(0)])).warnings == []


def test_review_warns_that_a_subscription_without_churn_is_always_beautiful():
    """Ноль оттока допустим, но не нейтрален: база тогда только растёт, и подписка
    выходит красивой при любых прочих допущениях."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(20)] * 12, churn_monthly=D(0))
    model = _model(sub, n=12)
    review = run_review(ReviewContext(model=model, result=run(model)))
    found = [f for f in review.findings if f.id == "assumptions.subscription_without_churn"]
    assert len(found) == 1 and found[0].severity == "warning"
    assert "Подписка" in found[0].detail


def test_the_churn_warning_is_silent_when_churn_is_set():
    """Тишина правила — половина его смысла: сработавшее всегда, оно ничего не значит."""
    sub = Subscription(starting_base=D(100), new_per_month=[D(20)] * 12,
                       churn_monthly=D("0.02"))
    model = _model(sub, n=12)
    review = run_review(ReviewContext(model=model, result=run(model)))
    assert not [f for f in review.findings
                if f.id == "assumptions.subscription_without_churn"]
    # И у модели вовсе без подписки — тоже тишина.
    plain = _model(None, n=12, volume=[D(5)] * 12)
    plain_review = run_review(ReviewContext(model=plain, result=run(plain)))
    assert not [f for f in plain_review.findings
                if f.id == "assumptions.subscription_without_churn"]


# --- Один состав на три выхода (экран, документ, выгрузка) ---

def test_the_document_carries_the_base_and_the_churn_column():
    """База уходит в бизнес-план: раздел без выбытия отдельной колонкой читался бы как
    рост без потерь — по одному приросту одно от другого не отличить."""
    from datetime import date as _date
    from io import BytesIO

    from docx import Document

    from app.docgen import build_business_plan_docx
    from calc_core.review import run_review as _review
    from calc_core.review.opinion import build_opinion

    sub = Subscription(starting_base=D(100), new_per_month=[D(20)] * 6,
                       churn_monthly=D("0.05"))
    model = _model(sub, n=6)
    result = run(model)
    opinion = build_opinion(_review(ReviewContext(model=model, result=result)), result)
    doc = Document(BytesIO(build_business_plan_docx(
        model, result, opinion, project_name="Подписка", today=_date(2026, 7, 1))))
    text = "\n".join(p.text for p in doc.paragraphs)
    headers = [c.text for t in doc.tables for c in t.rows[0].cells]

    assert "Абонентская база" in text
    assert "Ушло всего" in headers
    assert "отток не задан в модели" in text          # ноль объяснён, а не оставлен нулём


def test_the_document_skips_the_section_without_subscriptions():
    """Раздел «Абонентская база» в документе без подписок — пустая обещанная таблица."""
    from datetime import date as _date
    from io import BytesIO

    from docx import Document

    from app.docgen import build_business_plan_docx
    from calc_core.review import run_review as _review
    from calc_core.review.opinion import build_opinion

    model = _model(None, n=3, volume=[D(5)] * 3)
    result = run(model)
    opinion = build_opinion(_review(ReviewContext(model=model, result=result)), result)
    doc = Document(BytesIO(build_business_plan_docx(
        model, result, opinion, project_name="Обычный", today=_date(2026, 7, 1))))
    assert "Абонентская база" not in "\n".join(p.text for p in doc.paragraphs)
