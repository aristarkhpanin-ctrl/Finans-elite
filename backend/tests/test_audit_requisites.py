"""Реквизиты документа и подписи (Финанс-Аудит; SPEC, Приложение Х).

Печатный бланк был отчётом о проверке, но не документом сделки: ни адресата, ни номера,
ни подписи. Пустые линии под выдуманными должностями печатать было нельзя — и правильный
выход не «нарисовать линии», а завести настоящих подписантов.

Проверяется то, на чём держится честность документа: подписи не бывает без имени,
неподписанный документ называет себя рабочим материалом, опечатка в ИНН называется, но
введённое значение не выбрасывается, а несверенность с реестром печатается всегда.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document

from app.audit_docgen import build_audit_docx
from audit_core import build_requisites, review_case
from audit_core.models import AuditSubjectModel
from audit_core.requisites import UNSIGNED, valid_inn, valid_ogrn

REQ = {
    "number": "ДД-14/2026", "date": "2026-09-05",
    "addressee": "Инвестиционному комитету ООО «Фонд»",
    "subject_full_name": "Общество с ограниченной ответственностью «Цель»",
    "subject_inn": "7707083893", "subject_ogrn": "1027700132195",
    "subject_address": "Москва, ул. Примерная, 1",
    "executor_name": "И. Петров", "executor_role": "Аналитик",
    "approver_name": "А. Сидорова", "approver_role": "Партнёр",
}


def model(**report) -> AuditSubjectModel:
    return AuditSubjectModel.model_validate({
        "name": "ООО «Цель»",
        "periods": [{"label": "2024", "kind": "year"}],
        "balance": {"A_CASH": ["100"], "P_EQUITY": ["100"]},
        "income": {"I_REVENUE": ["500"]},
        "report": report,
    })


# --- подписи ---

def test_unsigned_document_says_so_first():
    """Бланк со строгой вёрсткой и без подписи читается как заключение."""
    v = build_requisites(model())
    assert v.signed is False and v.signatures == []
    assert v.caveats[0] == UNSIGNED


def test_signature_needs_a_name():
    """Должность без имени — заготовка, а не подпись."""
    v = build_requisites(model(approver_role="Партнёр"))
    assert v.signed is False
    # Но и пропасть молча она не вправе: человек заполнил половину.
    assert any("без имени" in c for c in v.caveats)


def test_name_without_a_role_is_still_a_signature():
    # Должность необязательна: подписывает человек, а не должность.
    v = build_requisites(model(executor_name="И. Петров"))
    assert v.signed is True
    assert v.signatures[0].name == "И. Петров" and v.signatures[0].role == ""
    assert UNSIGNED not in v.caveats


def test_both_signatories_are_kept_in_order():
    v = build_requisites(model(**REQ))
    assert [(s.name, s.role) for s in v.signatures] == [
        ("И. Петров", "Аналитик"), ("А. Сидорова", "Партнёр")]


# --- реквизиты ---

def test_inn_typo_is_named_but_the_value_survives():
    """Платформа не вправе ни подтвердить ИНН, ни выбросить введённое."""
    v = build_requisites(model(subject_inn="7707083894", executor_name="И. Петров"))
    assert v.subject_inn == "7707083894"
    assert any("контрольной цифры" in c for c in v.caveats)


def test_correct_inn_and_ogrn_raise_no_caveats():
    v = build_requisites(model(**REQ))
    assert v.caveats == []
    assert v.subject_ogrn == "1027700132195"


def test_ogrn_typo_is_named():
    v = build_requisites(model(subject_ogrn="1027700132196", executor_name="И. Петров"))
    assert any("ОГРН" in c and "контрольной" in c for c in v.caveats)


def test_registry_check_is_named_as_not_done():
    """«Не сверено с ЕГРЮЛ» печатается всегда: реквизит с видом проверенного хуже
    отсутствующего."""
    v = build_requisites(model(**REQ))
    assert any("ЕГРЮЛ" in t for t in v.not_computed)
    assert any("полномочи" in t for t in v.not_computed)


def test_checksums_accept_real_numbers_and_reject_neighbours():
    # Самопроверка алгоритма: без неё «валидация» могла бы принимать что угодно.
    assert valid_inn("7707083893") and not valid_inn("7707083894")
    assert valid_inn("500100732259") and not valid_inn("500100732250")
    assert not valid_inn("123") and not valid_inn("77070838ab")
    assert valid_ogrn("1027700132195") and not valid_ogrn("1027700132196")
    assert valid_ogrn("304500116000157") and not valid_ogrn("304500116000158")


def test_empty_block_is_inert():
    """Дело без реквизитов ведёт себя как прежде — только честно неподписанное."""
    v = build_requisites(model())
    assert v.filled is False
    assert (v.number, v.addressee, v.subject_inn, v.date) == ("", "", "", None)


def test_own_date_replaces_the_formation_date():
    v = build_requisites(model(date="2026-09-05"))
    assert v.date == date(2026, 9, 5) and v.filled is True


# --- разбор дела и документ ---

def test_review_carries_requisites():
    r = review_case(model(**REQ), deep=False)
    assert r.requisites.signed is True
    assert r.requisites.addressee.startswith("Инвестиционному комитету")


def test_document_prints_addressee_and_signatures():
    blob = build_audit_docx(review_case(model(**REQ), deep=False),
                            subject_name="ООО «Цель»", today=date(2026, 9, 8))
    text = "\n".join(p.text for p in Document(BytesIO(blob)).paragraphs)
    assert "Инвестиционному комитету ООО «Фонд»" in text
    assert "ДД-14/2026" in text and "05.09.2026" in text
    assert "И. Петров" in text and "А. Сидорова" in text


def test_document_without_signatures_says_it_is_a_draft():
    blob = build_audit_docx(review_case(model(), deep=False),
                            subject_name="ООО «Цель»", today=date(2026, 9, 8))
    text = "\n".join(p.text for p in Document(BytesIO(blob)).paragraphs)
    assert "рабочий материал" in text
    # Дата документа не выдаётся за собственную: напечатана дата формирования.
    assert "дата формирования" in text


def test_analysis_response_carries_requisites(client, auth_headers):
    sid = client.post("/api/v1/audit/subjects",
                      json={"name": "ООО «Цель»",
                            "model": model(**REQ).model_dump(mode="json")},
                      headers=auth_headers).json()["id"]
    body = client.post(f"/api/v1/audit/subjects/{sid}/analyze",
                       headers=auth_headers).json()["requisites"]
    assert body["signed"] is True and body["subject_inn"] == "7707083893"
    assert [s["name"] for s in body["signatures"]] == ["И. Петров", "А. Сидорова"]
